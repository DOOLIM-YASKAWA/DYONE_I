"""사양정리 초안 pptx 뒤에 '부서별 확인' 장을 붙인다: 현황 1장 + 역할(부서)마다 한 장(항목이 많으면 이어지는 장).
- 내용 JSON: {title, roles: [{dept, group, budget, items: [{title, status, evidence, incident, note, by}]}]}
- 마지막 장이 Ending Page(빈 레이아웃)면 그 앞에 넣는다. 본문 장과 같은 레이아웃을 쓴다.
사용: python spec_dept_pages.py <초안.pptx> <내용.json> <출력.pptx>
"""
import sys, json
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt, Emu

in_path, json_path, out_path = sys.argv[1:4]
C = json.load(open(json_path, encoding="utf-8"))
prs = Presentation(in_path)
W, H = prs.slide_width, prs.slide_height
FONT = "맑은 고딕"
STATUS = ["미확인", "확인", "고객 질의 필요", "내부 결정 필요", "해당 없음"]
COLOR = {
    "미확인": RGBColor(0x6B, 0x72, 0x80), "확인": RGBColor(0x0B, 0x7A, 0x3E),
    "고객 질의 필요": RGBColor(0xC2, 0x41, 0x0C), "내부 결정 필요": RGBColor(0x1D, 0x4E, 0xD8),
    "해당 없음": RGBColor(0x9C, 0xA3, 0xAF),
}
HEAD_FILL = RGBColor(0x1F, 0x3A, 0x5F)
PER_SLIDE = 9

slides = list(prs.slides)
ending = slides[-1] if slides and len(slides[-1].shapes) == 0 else None
body_layout = slides[2].slide_layout if len(slides) > 2 else prs.slide_layouts[len(prs.slide_layouts) - 1]


def text(slide, x, y, w, h, s, size, bold=False, color=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
        if color: r.font.color.rgb = color
    return tb


def cell(c, s, size, bold=False, color=None, fill=None):
    c.text = ""
    tf = c.text_frame
    tf.word_wrap = True
    for i, line in enumerate((s or "").split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
        if color: r.font.color.rgb = color
    c.margin_left = c.margin_right = Emu(60000)
    c.margin_top = c.margin_bottom = Emu(30000)
    if fill is not None:
        c.fill.solid(); c.fill.fore_color.rgb = fill


def new_slide(title, sub):
    s = prs.slides.add_slide(body_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    m = int(W * 0.04)
    text(s, m, int(H * 0.03), W - 2 * m, int(H * 0.07), title, 24, bold=True)
    if sub:
        text(s, m, int(H * 0.10), W - 2 * m, int(H * 0.05), sub, 13, color=RGBColor(0x55, 0x5F, 0x6D))
    return s, m


def table(slide, m, rows, cols, widths):
    top = int(H * 0.16)
    shape = slide.shapes.add_table(rows, cols, m, top, W - 2 * m, int(H * 0.05) * rows)
    t = shape.table
    total = sum(widths)
    for i, w in enumerate(widths):
        t.columns[i].width = int((W - 2 * m) * w / total)
    return t


roles = C["roles"]

# ---------- 현황: 역할 × 상태 건수
s, m = new_slide("부서별 확인 현황", f'{C.get("title", "")} — 사양서 확인 항목 상태(역할별)')
t = table(s, m, len(roles) + 1, len(STATUS) + 2, [3, 2] + [1.3] * len(STATUS))
for j, h in enumerate(["부서(역할)", "구분"] + STATUS):
    cell(t.cell(0, j), h, 12, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), fill=HEAD_FILL)
for i, r in enumerate(roles, start=1):
    cell(t.cell(i, 0), r["dept"], 11, bold=True)
    cell(t.cell(i, 1), r.get("group") or "", 11)
    for j, st in enumerate(STATUS):
        n = sum(1 for it in r["items"] if it["status"] == st)
        cell(t.cell(i, j + 2), str(n) if n else "", 11, bold=n > 0 and st != "확인", color=COLOR[st])

# ---------- 역할마다 한 장(항목이 많으면 나눔)
for r in roles:
    items = r["items"]
    chunks = [items[i:i + PER_SLIDE] for i in range(0, len(items), PER_SLIDE)] or [[]]
    for k, chunk in enumerate(chunks):
        sub = f'{r.get("group") or ""}' + (f' · 실행예산: {r["budget"]}' if r.get("budget") else "")
        title = f'부서별 확인 — {r["dept"]}' + (f" ({k + 1}/{len(chunks)})" if len(chunks) > 1 else "")
        s, m = new_slide(title, sub)
        t = table(s, m, len(chunk) + 1, 4, [4, 1.6, 6, 3.4])
        for j, h in enumerate(["확인 항목", "상태", "사양서 근거", "코멘트 · 담당"]):
            cell(t.cell(0, j), h, 12, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), fill=HEAD_FILL)
        for i, it in enumerate(chunk, start=1):
            name = it["title"] + (f'\n⚠ {it["incident"]}' if it.get("incident") else "")
            cell(t.cell(i, 0), name, 11, bold=True)
            cell(t.cell(i, 1), it["status"], 11, bold=True, color=COLOR.get(it["status"]))
            cell(t.cell(i, 2), it.get("evidence") or "", 10)
            note = "\n".join(x for x in [it.get("note") or "", f'— {it["by"]}' if it.get("by") else ""] if x)
            cell(t.cell(i, 3), note, 10)

# Ending Page를 맨 뒤로
if ending is not None:
    lst = prs.slides._sldIdLst
    el = list(lst)[len(slides) - 1]
    lst.remove(el); lst.append(el)

prs.save(out_path)
