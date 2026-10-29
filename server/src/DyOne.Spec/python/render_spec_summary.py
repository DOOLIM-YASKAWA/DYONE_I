"""사양정리 초안 렌더러: 빈 양식 틀(pptx) + 내용 JSON → 사양정리 pptx.
- 양식(글꼴·위치·색·표 구조)은 빈 틀 그대로, 값만 JSON에서.
- JSON 표시: **굵게** {{파랑=사양서 값}} [[빨강=선정·확인 필요]] <<초록=기존 이력(번호 → 기존 이력 근거 장)>>.  칸·항목별 근거(사양서 쪽)는 슬라이드 노트로.
- 그림(레이아웃·일정표)은 사양서 PDF의 쪽·영역(pt)을 잘라 넣는다.
사용: python render_spec_summary.py <빈틀.pptx> <내용.json> <출력.pptx> <임시폴더>
"""
import sys, json, copy, math, os
import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt, Inches

tpl_path, json_path, out_path, tmp = sys.argv[1:5]
C = json.load(open(json_path, encoding="utf-8"))
prs = Presentation(tpl_path)
RED, BLUE, BLACK = RGBColor(0xFF, 0, 0), RGBColor(0, 0, 0xFF), RGBColor(0, 0, 0)
GREEN = RGBColor(0x00, 0x80, 0x3C)
MARKS = ("**", "[[", "]]", "{{", "}}", "<<", ">>")
pdf = fitz.open(C["source_pdf"]) if C.get("source_pdf") else None  # 사양서가 PDF가 아니면 그림을 잘라 붙이지 않는다


def crop(spec, name):
    path = os.path.join(tmp, name)
    x0, y0, x1, y1 = spec["clip"]
    pdf[spec["page"] - 1].get_pixmap(dpi=220, clip=fitz.Rect(x0, y0, x1, y1)).save(path)
    return path


def segments(text):
    out, i, bold, color, buf = [], 0, False, None, ""
    while i < len(text):
        two = text[i:i + 2]
        if two in MARKS:
            if buf: out.append((buf, bold, color)); buf = ""
            bold = (not bold) if two == "**" else bold
            color = {"[[": RED, "{{": BLUE, "<<": GREEN}.get(two, None if two in ("]]", "}}", ">>") else color)
            i += 2; continue
        buf += text[i]; i += 1
    if buf: out.append((buf, bold, color))
    return out


def fill(owner, text, size=None, keep=False):
    """keep=True: 틀의 글자 서식(굵기·색)을 그대로 둔다(제목·표 머리줄). 표시(**·{{·[[)가 있는 조각만 바꾼다."""
    tf = owner.text_frame
    p0 = tf.paragraphs[0]
    r0 = next((r for p in tf.paragraphs for r in p.runs), None)
    base = copy.deepcopy(r0._r.rPr) if r0 is not None and r0._r.rPr is not None else None
    ppr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(text.split("\n")):
        p = p0 if i == 0 else tf.add_paragraph()
        if i > 0 and ppr is not None:
            p._p.insert(0, copy.deepcopy(ppr))
        for seg, bold, color in segments(line) or [(" ", False, None)]:
            r = p.add_run(); r.text = seg
            if base is not None:
                r._r.insert(0, copy.deepcopy(base))
            if size: r.font.size = Pt(size)
            if keep:
                if bold: r.font.bold = True
                if color: r.font.color.rgb = color
            else:
                r.font.bold = bold
                r.font.color.rgb = color or BLACK


def plain(s):
    for t in MARKS:
        s = s.replace(t, "")
    return s


def norm(s):
    return "".join(ch for ch in s.lower() if ch.isalnum())


def set_titles(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and "{고객}" in sh.text_frame.text:
            fill(sh, C["title"], keep=True)


# ---------- 1 표지 ----------
for sh in prs.slides[0].shapes:
    if sh.has_text_frame and "{고객}" in sh.text_frame.text:
        fill(sh, C["cover"], keep=True)

# ---------- 2 요약표 ----------
s2 = prs.slides[1]
set_titles(s2)
notes = []
def split_value_column(table, n):
    """값 열(1번)을 공정 수(n)만큼 나눈다: 열 너비를 n등분하고 모든 줄에 칸을 복제(영업부 양식 FRT | RR)."""
    from pptx.oxml.ns import qn
    tbl = table._tbl
    cols = tbl.tblGrid.findall(qn("a:gridCol"))
    w = int(cols[1].get("w"))
    # 공정이 많으면(여러 공정 프로젝트, 화성4 7개) 기타 사항 열을 줄여 공정 열에 나눠 준다
    if n >= 4 and len(cols) >= 3:
        nw = int(cols[-1].get("w"))
        give = int(nw * 0.45)
        cols[-1].set("w", str(nw - give))
        w += give
    cols[1].set("w", str(w // n))
    for _ in range(n - 1):
        cols[1].addnext(copy.deepcopy(cols[1]))
    for tr in tbl.tr_lst:
        tcs = tr.tc_lst
        for _ in range(n - 1):
            tcs[1].addnext(copy.deepcopy(tcs[1]))


def add_missing_rows(table):
    """템플릿 표를 초안 줄에 맞춘다: 초안에 없는 템플릿 줄(예: Power Unit — 2026-10-27 표에서 뺌)은 지우고,
    템플릿에 없는 줄(예: Others — 영업부 사양정리는 표 맨 아래)은 마지막 줄을 복사해 붙인다."""
    keys = [norm(k) for k in C["rows"]]
    for i, r in reversed(list(enumerate(table.rows))):
        h = norm(r.cells[0].text)
        if i > 0 and h and not any(k.startswith(h) or h.startswith(k) for k in keys):
            tr = table._tbl.tr_lst[i]
            tr.getparent().remove(tr)
    heads = [norm(r.cells[0].text) for r in table.rows]
    total = sum(r.height for r in table.rows)
    added = False
    for key in C["rows"]:
        if any(h and (h.startswith(norm(key)) or norm(key).startswith(h)) for h in heads):
            continue
        tr = table._tbl.tr_lst[-1]
        tr.addnext(copy.deepcopy(tr))
        fill(table.rows[len(table.rows) - 1].cells[0], key, keep=True)
        heads.append(norm(key))
        added = True
    if added:  # 줄이 늘어도 슬라이드 밖으로 넘치지 않게: 머리줄 외 최소 높이를 낮춰 글 양만큼만 차지하게 한다
        rows = list(table.rows)[1:]
        low = min(int(total / len(table.rows) * 0.75), 288000)  # 0.8cm 이하
        for r in rows:
            r.height = low


# 표 글자 크기: 17줄(+Others)이 슬라이드 한 장에 들어가도록 14 → 12pt
CELL_PT = 12 if len(C.get("columns") or []) < 4 else 9  # 공정 열이 많으면 글자를 줄여 표가 한 장에 들어가게
columns = C.get("columns") or []
ncol = max(1, len(columns))
for sh in list(s2.shapes):
    if sh.has_table:
        table = sh.table
        add_missing_rows(table)
        if ncol > 1:
            split_value_column(table, ncol)
        for i, row in enumerate(table.rows):
            note_cell = row.cells[ncol + 1]
            if i == 0:
                if ncol > 1:
                    for j, title in enumerate(columns):
                        fill(row.cells[1 + j], title, 12, keep=True)
                else:
                    fill(row.cells[1], C["col_header"], keep=True)
                continue
            head = norm(row.cells[0].text)
            key = next((k for k in C["rows"] if head.startswith(norm(k)) or norm(k).startswith(head)), None)
            rec = (C["rows"][key] + ["", "", None])[:4] if key else ["[[사양서에 없음 — 선정 필요]]", " ", "", None]
            val, note, src, spans = rec[0], rec[1], rec[2], rec[3] if isinstance(rec[3], list) else None
            if isinstance(val, list) and ncol > 1:  # 공정별 값: 칸마다(화면에서 합친 칸은 spans대로 병합)
                for j in range(ncol):
                    span = spans[j] if spans and j < len(spans) else 1
                    if span == 0:
                        continue
                    if span > 1:
                        row.cells[1 + j].merge(row.cells[min(ncol, j + span)])
                    fill(row.cells[1 + j], val[j] if j < len(val) and val[j] else " ", CELL_PT)
            else:  # 모든 공정에 같은 값: 값 칸을 한 칸으로 합친다
                if ncol > 1:
                    row.cells[1].merge(row.cells[ncol])
                fill(row.cells[1], "\n".join(val) if isinstance(val, list) else val, CELL_PT)
            fill(note_cell, note or " ", CELL_PT)
            if key: notes.append(f"- {key}: {src}")
    elif sh.has_text_frame and "Total" in sh.text_frame.text:
        fill(sh, C["total"], keep=True)
    elif sh.has_text_frame and "LINE" in sh.text_frame.text:
        if C.get("lines"): fill(sh, C["lines"])
        else: sh._element.getparent().remove(sh._element)
if C.get("layout_image"):
    s2.shapes.add_picture(crop(C["layout_image"], "layout.png"), Inches(5.32), Inches(0.84), height=Inches(1.3))
if C.get("history"):  # 범례: 글자 색이 무엇을 뜻하는지
    lg = s2.shapes.add_textbox(Inches(11.6), Inches(1.75), Inches(6.4), Inches(0.5))
    lg.text_frame.word_wrap = True
    fill(lg, "{{파랑}} 사양서 값 · [[빨강]] 선정·확인 필요 · 모델 근거는 '모델 제안 근거' 장", 12)
s2.notes_slide.notes_text_frame.text = "칸별 근거 (사양서 쪽)\n" + "\n".join(notes)

# ---------- 3 일정 ----------
s3 = prs.slides[2]
set_titles(s3)
edu_box = note_box = None
for sh in s3.shapes:
    if sh.has_text_frame and "{교육" in sh.text_frame.text: edu_box = sh
    if sh.has_text_frame and "{일정" in sh.text_frame.text: note_box = sh
class _Box:  # 그림 대신 간트 표를 그렸을 때 위치 계산용
    def __init__(self, left, top, width, height):
        self.left, self.top, self.width, self.height = left, top, width, height


if "schedule_image" not in C and "schedule_gantt" not in C:  # 일정표를 못 찾음: 요약 한 줄만
    si = {"left": 0.99, "top": 2.1, "page": "-"}
    pic = _Box(Inches(0.99), Inches(2.1), Inches(8.0), Inches(0.1))
elif "schedule_image" in C:  # 사양서에 일정표가 있으면 그대로 붙인다
    si = C["schedule_image"]
    pic = s3.shapes.add_picture(crop(si, "schedule.png"), Inches(si["left"]), Inches(si["top"]),
                                **({"height": Inches(si["height"])} if "height" in si else {"width": Inches(si["width"])}))
else:  # 없으면 사양서에 나온 날짜로 달력(간트)을 그린다
    g = C["schedule_gantt"]
    si = {"left": g.get("left", 0.99), "top": g.get("top", 2.1), "page": g.get("src", "-")}
    months, grows = g["months"], g["rows"]
    name_w, cell_w = g.get("name_width", 4.2), g.get("cell_width", 0.95)
    gt = s3.shapes.add_table(len(grows) + 1, len(months) + 1, Inches(si["left"]), Inches(si["top"]),
                             Inches(name_w + cell_w * len(months)), Inches(0.42) * (len(grows) + 1)).table
    gt.columns[0].width = Inches(name_w)
    for i in range(len(months)):
        gt.columns[i + 1].width = Inches(cell_w)
    for c, m in enumerate(["구분"] + months):
        cell = gt.cell(0, c); cell.text = m
        p = cell.text_frame.paragraphs[0]; p.alignment = 2
        p.runs[0].font.size = Pt(13); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0x00, 0x2F, 0x87)
    for r, (name, marks) in enumerate(grows, start=1):
        cell = gt.cell(r, 0); cell.text = ""
        fill(cell, name, 12)
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xD9, 0xDF, 0xEA)
        for c in range(len(months)):
            cell = gt.cell(r, c + 1)
            mark = marks.get(str(c))
            cell.text = ""
            cell.fill.solid()
            if mark:
                color, label = (mark + "|").split("|")[:2]
                cell.fill.fore_color.rgb = RGBColor.from_string(color)
                if label:
                    cell.text = label
                    p = cell.text_frame.paragraphs[0]; p.alignment = 2
                    p.runs[0].font.size = Pt(10); p.runs[0].font.bold = True
            else:
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    pic = _Box(Inches(si["left"]), Inches(si["top"]), Inches(name_w + cell_w * len(months)), Inches(0.42) * (len(grows) + 1))
et = C.get("edu_table") or {"left": 9.4, "top": 3.6, "widths": [], "rows": []}
right = pic.left + pic.width + Inches(0.3)
side = Inches(et["left"]) >= right - Inches(0.05)  # 일정표 오른쪽에 둘지(세로로 긴 일정표), 아래에 둘지
if note_box is not None:
    fill(note_box, C["schedule_note"], 18, keep=True)
    note_box.text_frame.word_wrap = True
    if side:
        note_box.left, note_box.top, note_box.width = Inches(et["left"]), Inches(si["top"]), Inches(9.8)
    else:
        note_box.left, note_box.top, note_box.width = Inches(si["left"]), pic.top + pic.height + Inches(0.1), Inches(16.5)
rows = et["rows"]
if edu_box is not None:
    if rows:
        fill(edu_box, C["edu_title"], keep=True)
        edu_box.left, edu_box.top = Inches(et["left"]), Inches(et["top"]) - Inches(0.55)
    else:  # 교육 일정표가 없으면 제목도 지운다
        edu_box._element.getparent().remove(edu_box._element)
if rows:
    t = s3.shapes.add_table(len(rows), len(rows[0]), Inches(et["left"]), Inches(et["top"]),
                            Inches(sum(et["widths"])), Inches(0.55) * len(rows)).table
    for i, w in enumerate(et["widths"][:len(rows[0])]):
        t.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = t.cell(r, c); cell.text = str(v) or " "
            para = cell.text_frame.paragraphs[0]; para.alignment = 2
            para.runs[0].font.size = Pt(14); para.runs[0].font.color.rgb = BLACK
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0x99) if r == 0 else RGBColor(0xFF, 0xFF, 0xFF)
    if C.get("edu_caption"):
        cap = s3.shapes.add_textbox(Inches(et["left"]), Inches(et["top"]) + Inches(0.55) * len(rows) + Inches(0.2),
                                    Inches(sum(et["widths"]) + 0.5), Inches(0.6))
        cap.text_frame.word_wrap = True
        fill(cap, C["edu_caption"], 12)
s3.notes_slide.notes_text_frame.text = (f"근거: 일정표 사양서 p{si['page']} (그대로 붙임)" if "schedule_image" in C
                                        else f"사양서에 일정표 없음 → 사양서 날짜로 달력 작성. 근거: {C['schedule_gantt'].get('src', '')}" if "schedule_gantt" in C
                                        else "일정표를 찾지 못함 — 사양서 확인")


# ---------- 4·5 체크리스트 ----------
def checklist(slide, items, size, top, step, chars, title_note):
    set_titles(slide)
    proto = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip().startswith("□"))
    el0 = proto._element
    y = top
    for text, _src in items:
        el = copy.deepcopy(el0)
        slide.shapes._spTree.append(el)
        sh = slide.shapes[-1]
        fill(sh, "□ " + text, size)
        sh.left, sh.top, sh.width = Inches(1.38), Inches(y), Inches(16.4)
        sh.text_frame.word_wrap = True
        y += step * max(1, math.ceil(len(plain(text)) / chars)) + 0.02
    el0.getparent().remove(el0)
    slide.notes_slide.notes_text_frame.text = title_note + "\n" + "\n".join(f"- {plain(t)[:40]} … : {s}" for t, s in items)


n_etc = len(C["etc"])
etc_size = 14 if n_etc > 20 else 16
checklist(prs.slides[3], C["etc"], etc_size, 1.2, 0.33 if etc_size == 14 else 0.38, 110 if etc_size == 14 else 95, "항목별 근거 (사양서 쪽)")
# 제공 항목 장: 템플릿 PCS 장을 복사해 소제목만 바꾼다(사양서에 제공 요구가 있을 때만) — PCS 체크리스트를 채우기 전에 복사(□ 틀이 남아 있게)
supply_slide = None
if C.get("supply"):
    src5 = prs.slides[4]
    supply_slide = prs.slides.add_slide(src5.slide_layout)
    for ph in list(supply_slide.placeholders):
        ph._element.getparent().remove(ph._element)
    for sh in src5.shapes:
        supply_slide.shapes._spTree.append(copy.deepcopy(sh._element))
    for sh in supply_slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "PCS 사항":
            fill(sh, "제공 항목", keep=True)
checklist(prs.slides[4], C["pcs"], 18, 2.3, 0.42, 80, "항목별 근거 (사양서 쪽)")
if supply_slide is not None:
    checklist(supply_slide, C["supply"], 18, 2.3, 0.42, 80, "항목별 근거 (사양서 쪽)")


# ---------- 기존 이력 근거 (사양서의 "기설치·기존과 동일"을 과거 프로젝트에서 찾은 근거) ----------
def move_slide(prs, old, new):
    lst = prs.slides._sldIdLst
    el = list(lst)[old]
    lst.remove(el); lst.insert(new, el)


if supply_slide is not None:  # 제공 항목 장은 PCS 장 바로 뒤로
    move_slide(prs, len(prs.slides) - 1, 5)


H = C.get("history")
if H:
    src = prs.slides[3]
    sl = prs.slides.add_slide(src.slide_layout)
    for ph in list(sl.placeholders):
        ph._element.getparent().remove(ph._element)
    for sh in list(src.shapes)[:2]:  # 제목 + 소제목 틀 복사
        sl.shapes._spTree.append(copy.deepcopy(sh._element))
    fill(sl.shapes[1], H["title"], keep=True)
    sl.shapes[1].width = Inches(6)
    intro = sl.shapes.add_textbox(Inches(0.99), Inches(1.25), Inches(16.8), Inches(0.5))
    intro.text_frame.word_wrap = True
    fill(intro, H["intro"], 13)
    # 구분: 사양서가 '기설치와 동일'을 요구한 항목인지, 모델이 없어 제안하는 항목인지
    cols = [("번호", 0.6), ("구분", 1.5), ("항목", 1.6), ("제안 모델", 3.3), ("근거 (연도·공정·프로젝트)", 3.3), (r"출처 파일 (Q:\01_프로젝트\…)", 5.1), ("확인", 1.6)]
    rows = H["rows"]
    t = sl.shapes.add_table(len(rows) + 1, len(cols), Inches(0.99), Inches(1.95),
                            Inches(sum(w for _, w in cols)), Inches(0.4) * (len(rows) + 1)).table
    for c, (name, w) in enumerate(cols):
        t.columns[c].width = Inches(w)
        cell = t.cell(0, c); cell.text = name
        r0 = cell.text_frame.paragraphs[0].runs[0]
        r0.font.size = Pt(12); r0.font.bold = True; r0.font.color.rgb = BLACK
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0x99)
    for r, h in enumerate(rows, start=1):
        vals = [f"<<{h['id']}>>", h.get("kind", ""), h["item"], f"<<{h['value']}>>", h["basis"], h["file"], h.get("check", "")]
        for c, v in enumerate(vals):
            cell = t.cell(r, c); cell.text = ""
            fill(cell, v, 9 if c == 5 else 11)
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if H.get("note"):
        nb = sl.shapes.add_textbox(Inches(0.99), Inches(9.6), Inches(16.8), Inches(0.5))
        nb.text_frame.word_wrap = True
        fill(nb, H["note"], 11)
    sl.notes_slide.notes_text_frame.text = "기존 이력 근거 – 출처 전체 경로\n" + "\n".join(
        f"{h['id']} Q:\\01_프로젝트\\{h['file']}" for h in rows)
    move_slide(prs, len(prs.slides) - 1, len(prs.slides) - 2)  # 끝 장 앞으로

prs.save(out_path)
print("saved", out_path)
